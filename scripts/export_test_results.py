"""Run the Tripzy test suite and save the results to an Excel workbook.

Usage:
    python scripts/export_test_results.py --tester "Your Name"
    python scripts/export_test_results.py --skip-e2e        # without browser tests

Output: reports/tripzy_test_results.xlsx
"""

import argparse
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "reports" / "tripzy_test_results.xlsx"

# PBL Assignment 7 cases (spec §11.3). "verified" describes what the automated test checks,
# and is used as the Actual Result only when that test passed.
ACCEPTANCE = [
    ("BS-01", "Budget Splitter", "Add expense with valid amount, payer and participant",
     "Open trip → Budget Splitter → Add Expense → enter amount → choose payer and participant(s) → choose split option → save",
     "Expense is added successfully", "test_add_expense_valid",
     "Expense of ₹900.00 saved (HTTP 201) with the chosen payer; shares sum to the total; expense appears in the list."),
    ("BS-02", "Budget Splitter", "Add a negative expense", "Enter a negative amount and attempt to save",
     'Error message: "Enter a valid amount"', "test_negative_expense_rejected",
     'Save rejected (HTTP 422) with message "Enter a valid amount"; no expense was created.'),
    ("BS-03", "Budget Splitter", "Split equally among all members",
     "Add expense, select all trip members, choose equal split, save",
     "Expense is split among all selected members", "test_equal_split_all_members",
     "₹900.00 split into ₹300.00 for each of the 3 members."),
    ("BS-04", "Budget Splitter", "Split among selected members",
     "Add expense and select only required participants",
     "Only selected participants receive shares", "test_split_selected_members_only",
     "₹600.00 split ₹300.00 each between the 2 selected members; the unselected member has no share and a zero balance."),
    ("BS-05", "Budget Splitter", "Customise shares",
     "Enter details, payer and participants; set custom shares; verify total; save",
     "Expense is added and shares/amounts update correctly", "test_custom_split_reconciles_total",
     "₹500.00 saved with custom shares ₹100 / ₹250 / ₹150; net balances updated to +₹400 / −₹250 / −₹150."),
    ("BS-06", "Budget Splitter", "Delete an expense", "Locate expense → Delete → confirm if prompted",
     "Expense is deleted and balances update", "test_delete_expense_updates_balances",
     "Expense and its shares removed from the database; balances recalculated from the remaining expense."),
    ("TM-01", "Trip Management", "Create trip",
     "Open Trip Management → Create Trip → enter name, destination/details and optional dates → save",
     "New trip is created successfully", "test_create_trip",
     "Trip “Goa Weekend” created (HTTP 201) with destination and dates, and listed under the user's trips."),
    ("TM-02", "Trip Management", "Add members", "Open trip → Members → Add Member → select user → confirm",
     "Selected member is added", "test_owner_adds_trip_member",
     "User added as a trip member (HTTP 201) and can now open the trip."),
    ("TM-03", "Trip Management", "Remove member",
     "Open trip → Members → select member → Remove Member → confirm",
     "Selected member is removed", "test_owner_removes_trip_member",
     "Selected member removed (HTTP 204); other members unaffected."),
    ("TM-04", "Trip Management", "Delete trip", "Open trip settings → Delete Trip → confirm",
     "Selected trip and dependent data are deleted", "test_delete_trip_cascades_data",
     "Trip, its members, expenses, shares and saved locations deleted; another trip's data untouched."),
    ("LS-01", "Location Suggestion", "Add location",
     "Open trip → Locations → Add Location → search/select a place → Save",
     "Location is saved to the trip", "test_save_location_to_trip",
     "Selected place saved (HTTP 201) and visible to other trip members."),
    ("LS-02", "Location Suggestion", "Suggest popular locations", "Open Location Suggestions and load suggestions",
     "Popular places are shown", "test_popular_location_suggestions",
     "Suggestions for the trip destination returned, each with a name and category; Locations page loads."),
    ("LS-03", "Location Suggestion", "Enable location access",
     "Select Use My Location and allow the browser prompt",
     "App obtains location and uses it for nearby suggestions", "test_geolocation_granted_uses_location",
     "Browser test (Chromium): explanation shown first; after allowing, results headed “Popular places near you” for the device coordinates."),
    ("LS-04", "Location Suggestion", "Location permission denied",
     "Trigger location access and deny the browser prompt",
     "Permission-denied message is shown", "test_geolocation_denied_shows_message",
     "Browser test (Chromium): “Location permission denied. You can still search using your trip destination.” shown; no second prompt; destination search still works."),
]

AREAS = {
    "test_auth": "Authentication",
    "test_friends": "Friends",
    "test_friend_groups": "Friend groups",
    "test_trips": "Trips",
    "test_expenses": "Budget / expenses",
    "test_locations": "Locations",
    "test_acceptance_cases": "Acceptance (PBL 7)",
    "test_browser": "Browser end-to-end",
}

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
STATUS_FILLS = {
    "Pass": PatternFill("solid", fgColor="E3F6EC"),
    "Fail": PatternFill("solid", fgColor="FDECEA"),
    "Not run": PatternFill("solid", fgColor="FFF5D6"),
    "Skipped": PatternFill("solid", fgColor="FFF5D6"),
}
THIN = Side(style="thin", color="D8DFE3")


def run_pytest(skip_e2e: bool) -> tuple[Path, datetime, str]:
    xml_path = Path(tempfile.mkdtemp(prefix="tripzy-junit-")) / "results.xml"
    cmd = [sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider", f"--junitxml={xml_path}"]
    if skip_e2e:
        cmd += ["--ignore=tests/e2e"]
    started = datetime.now()
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    summary = (proc.stdout.strip().splitlines() or ["(no output)"])[-1]
    print(summary)
    if not xml_path.exists():
        sys.exit(f"pytest did not produce results:\n{proc.stdout}\n{proc.stderr}")
    return xml_path, started, summary


def parse_results(xml_path: Path) -> list[dict]:
    results = []
    for case in ET.parse(xml_path).getroot().iter("testcase"):
        module = case.get("classname", "").split(".")[-1]
        name = case.get("name", "")
        status, message = "Pass", ""
        for tag, label in (("failure", "Fail"), ("error", "Fail"), ("skipped", "Skipped")):
            el = case.find(tag)
            if el is not None:
                status = label
                message = (el.get("message") or el.text or "").strip().splitlines()[0][:500] if (el.get("message") or el.text) else ""
                break
        results.append({
            "file": case.get("classname", "").replace(".", "/") + ".py",
            "module": module,
            "name": name,
            "base_name": re.sub(r"\[.*\]$", "", name),
            "area": AREAS.get(module, module),
            "status": status,
            "seconds": float(case.get("time") or 0),
            "message": message,
        })
    return results


def style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def style_body(ws, first_row: int, last_row: int, ncols: int, status_col: int) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=ncols):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        status = row[status_col - 1]
        if status.value in STATUS_FILLS:
            status.fill = STATUS_FILLS[status.value]
            status.font = Font(name=FONT, size=10, bold=True)


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(results: list[dict], started: datetime, tester: str, summary: str, out: Path) -> None:
    by_name: dict[str, dict] = {}
    for r in results:
        by_name.setdefault(r["base_name"], r)

    wb = Workbook()

    # ---- Sheet 1: acceptance cases in the PBL spreadsheet layout ----
    ws = wb.active
    ws.title = "Acceptance (PBL 7)"
    headers = ["Test ID", "Module", "Test Scenario", "Test Steps / Input", "Expected Result", "Actual Result",
               "Status", "Executed Time", "Executed By", "Executed Date", "Automated Test"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for test_id, module, scenario, steps, expected, test_name, verified in ACCEPTANCE:
        r = by_name.get(test_name)
        if r is None:
            actual, status = "Not executed in this run (browser tests skipped).", "Not run"
            time_txt = date_txt = by = ""
        else:
            status = r["status"]
            actual = verified if status == "Pass" else f"{status}: {r['message'] or 'see pytest output'}"
            time_txt, date_txt, by = started.strftime("%H:%M:%S"), started.strftime("%d-%m-%Y"), tester
        ws.append([test_id, module, scenario, steps, expected, actual, status, time_txt, by, date_txt,
                   r["file"] + "::" + r["name"] if r else test_name])
    style_body(ws, 2, ws.max_row, len(headers), status_col=7)
    set_widths(ws, [9, 18, 26, 42, 30, 52, 10, 12, 20, 13, 48])
    ws.freeze_panes = "B2"
    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1,
            value=f"Executed Time / Date = when this automated run started. Executed By = value passed with --tester. "
                  f"Actual Result describes what the automated test verified; it is shown only when the test passed. "
                  f"pytest summary: {summary}").font = Font(name=FONT, size=9, italic=True)

    # ---- Sheet 2: every test ----
    all_ws = wb.create_sheet("All tests")
    all_headers = ["#", "Area", "Test file", "Test name", "Status", "Duration (s)", "Failure / skip message"]
    all_ws.append(all_headers)
    style_header(all_ws, 1, len(all_headers))
    for i, r in enumerate(results, start=1):
        all_ws.append([i, r["area"], r["file"], r["name"], r["status"], round(r["seconds"], 3), r["message"]])
        all_ws.cell(row=i + 1, column=6).number_format = "0.000"
    last = all_ws.max_row
    style_body(all_ws, 2, last, len(all_headers), status_col=5)
    set_widths(all_ws, [6, 20, 34, 62, 10, 13, 50])
    all_ws.freeze_panes = "A2"
    table = Table(displayName="AllTests", ref=f"A1:{get_column_letter(len(all_headers))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    all_ws.add_table(table)

    # ---- Sheet 3: summary (formulas over the All tests sheet) ----
    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "Tripzy — automated test run"
    sm["A1"].font = Font(name=FONT, bold=True, size=14)
    info = [("Run started", started.strftime("%d-%m-%Y %H:%M:%S")), ("Executed by", tester),
            ("Command", "python scripts/export_test_results.py"), ("pytest summary", summary)]
    for i, (k, v) in enumerate(info, start=3):
        sm.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True)
        sm.cell(row=i, column=2, value=v).font = Font(name=FONT)

    rng_status = f"'All tests'!$E$2:$E${last}"
    rng_area = f"'All tests'!$B$2:$B${last}"
    rng_time = f"'All tests'!$F$2:$F${last}"
    top = 9
    sm.cell(row=top, column=1, value="Area")
    for j, h in enumerate(["Total", "Pass", "Fail", "Skipped", "Pass rate", "Duration (s)"], start=2):
        sm.cell(row=top, column=j, value=h)
    style_header(sm, top, 7)
    areas = list(dict.fromkeys(r["area"] for r in results))
    for i, area in enumerate(areas, start=top + 1):
        sm.cell(row=i, column=1, value=area)
        sm.cell(row=i, column=2, value=f'=COUNTIF({rng_area},A{i})')
        sm.cell(row=i, column=3, value=f'=COUNTIFS({rng_area},A{i},{rng_status},"Pass")')
        sm.cell(row=i, column=4, value=f'=COUNTIFS({rng_area},A{i},{rng_status},"Fail")')
        sm.cell(row=i, column=5, value=f'=COUNTIFS({rng_area},A{i},{rng_status},"Skipped")')
        sm.cell(row=i, column=6, value=f'=IF(B{i}=0,0,C{i}/B{i})')
        sm.cell(row=i, column=7, value=f'=SUMIF({rng_area},A{i},{rng_time})')
    total = top + 1 + len(areas)
    sm.cell(row=total, column=1, value="Total")
    for col in "BCDEG":
        sm[f"{col}{total}"] = f"=SUM({col}{top + 1}:{col}{total - 1})"
    sm[f"F{total}"] = f"=IF(B{total}=0,0,C{total}/B{total})"
    for row in sm.iter_rows(min_row=top + 1, max_row=total, max_col=7):
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.row == total)
            cell.border = Border(bottom=THIN)
        row[5].number_format = "0.0%"
        row[6].number_format = "0.00"

    acc_row = total + 2
    sm.cell(row=acc_row, column=1, value="Acceptance cases (PBL 7)").font = Font(name=FONT, bold=True)
    acc_status = f"'Acceptance (PBL 7)'!$G$2:$G${len(ACCEPTANCE) + 1}"
    for k, (label, formula) in enumerate([
        ("Passed", f'=COUNTIF({acc_status},"Pass")'),
        ("Failed", f'=COUNTIF({acc_status},"Fail")'),
        ("Not run", f'=COUNTIF({acc_status},"Not run")'),
    ], start=acc_row + 1):
        sm.cell(row=k, column=1, value=label).font = Font(name=FONT)
        sm.cell(row=k, column=2, value=formula).font = Font(name=FONT)
    set_widths(sm, [26, 44, 10, 10, 10, 11, 13])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--tester", default="Automated (pytest)", help="Name for the 'Executed By' column")
    parser.add_argument("--skip-e2e", action="store_true", help="Skip the Playwright browser tests")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    xml_path, started, summary = run_pytest(args.skip_e2e)
    results = parse_results(xml_path)
    build_workbook(results, started, args.tester, summary, args.out)
    print(f"Saved {len(results)} results to {args.out}")


if __name__ == "__main__":
    main()
